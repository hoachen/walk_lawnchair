from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Launcher_Tool_App_Integration_Guide.md"
OUTPUT = ROOT / "Walk_Lawnchair_工具App接入与页面定制指南.docx"


def set_font(run, size=10.5, bold=None, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(paragraph, text, size=10.5):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size, color=(31, 78, 120))
        elif part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F7")
    set_cell_margins(cell, 100, 140, 100, 140)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_font(run, 8.5, color=(31, 78, 120))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    widths = [1.25, 2.25, 3.0]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.width = Inches(widths[col_index] if col_index < len(widths) else 2.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, 9.0)
            for run in p.runs:
                if row_index == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for level, size, color in ((1, 16, (46, 116, 181)), (2, 13, (46, 116, 181)), (3, 11.5, (31, 78, 120))):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Walk Lawnchair：工具 App 接入与页面定制指南")
    set_font(run, 22, bold=True, color=(11, 37, 69))
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("适用发布物：launcher-sdk 1.0.15  |  最后更新：2026-08-01")
    set_font(run, 10, color=(89, 89, 89))

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# ") or line.startswith("**适用") or line.startswith("**最后更新"):
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(doc, code)
        elif re.match(r"^\|.*\|$", line):
            table_lines = []
            while i < len(lines) and re.match(r"^\|.*\|$", lines[i]):
                if not re.match(r"^\|\s*[-:]+", lines[i]):
                    table_lines.append([item.strip() for item in lines[i].strip("|").split("|")])
                i += 1
            if table_lines:
                add_table(doc, table_lines)
            i -= 1
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 3")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.1)
            add_text(p, line[2:], 10)
        elif re.match(r"^[-*] \[[ x]\] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_text(p, line[2:], 10)
        elif re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_text(p, line[2:], 10)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_text(p, re.sub(r"^\d+\. ", "", line), 10)
        elif line.strip():
            p = doc.add_paragraph()
            add_text(p, line, 10.5)
        i += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Walk Lawnchair 集成文档")
    set_font(footer_run, 8.5, color=(89, 89, 89))
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
